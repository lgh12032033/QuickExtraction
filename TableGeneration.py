#!/usr/bin/env python
# -*- coding:utf-8 -*-
# @FileName  :TableGeneration.py
# @Time      :2021/8/9 11:48
# @Author    : 刘广鸿
from ScanfDoc import *
import xlwt
from config import *
from Project import *
import docx


class TableGeneration:

    def __init__(self, fileName):
        self.fileName = fileName
        self.work_book = xlwt.Workbook()  # 创建工作簿
        self.work_sheet = self.work_book.add_sheet("Sheet1")  # 新建表格
        self.cols = []
        self.__init_col()

    def __init_col(self):
        for column in config["columns"]:
            # columnName = list(column.keys())[0]
            if column["rule_Type"] == "reregular_noe":  # 单匹配
                Re_MatchRule = ReMatchRule(re.search)
            elif column["rule_Type"] == "reregular":
                Re_MatchRule = ReMatchRule()
            self.cols.append(Project(Re_MatchRule, column))
        self.__setTitle()

    def saveFile(self):
        self.work_book.save('%s.xls' % self.fileName)

    def __setSize(self):
        len = 625
        for item in config["columns"]:
            self.work_sheet.width = 36 * len

    def __setTitle(self):
        self.__setSize()
        for col in self.cols:
            col.init_colName(self.work_sheet)

    def writeProject(self, text, file):
        for poj in self.cols:
            poj.add_item(self, poj.getColumn()["lists_rule"], text)
        self.work_sheet.write(self.cols[0].getRow(), len(self.cols), file.split("\\")[-1])


if __name__ == '__main__':
    scan = ScanFile()
    filelist = scan.scanDocx()
    print("当前文件夹docx文件个数:%s" % len(filelist))  # 文件个数
    excel = TableGeneration(config["outputPath"])
    for filename in filelist:
        # print(filename)
        doc = docx.Document(filename)
        text = ""
        for para in doc.paragraphs:
            text += para.text
        excel.writeProject(text, filename)
    excel.saveFile()
    for doc in doc_files:
        os.rename(doc, doc[:-1])
